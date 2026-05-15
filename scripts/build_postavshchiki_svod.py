# -*- coding: utf-8 -*-
"""
Сводная ведомость папки docs/Поставщики.

Колонки: № папки, название папки в архиве, поставщик, вид деятельности / предмет договора.

Запуск:
  python scripts/build_postavshchiki_svod.py
"""

from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from delo_github_policy import (
    copy_register_to_case,
    ensure_registers_dir,
    update_register_opis_row,
    REGISTER_FILES,
)

BASE = Path(__file__).resolve().parent.parent
SUPPLIERS_ROOT = BASE / "docs" / "Поставщики"
APP_ID = "11.2"
KATALOG_PATH = BASE / "Дело3а-78-2026Таганай" / "Каталог_приложений.md"

NUM_RE = re.compile(r"^(\d+)\.\s*(.+)$", re.U)
PARENS_RE = re.compile(r"\(([^)]+)\)\s*$", re.U)
LEGAL_RE = re.compile(
    r"^(ИП|ООО|ОАО|ЗАО|ПАО|АО|САО|СТО|ГУП|МУП|ФГУП)\b",
    re.I | re.U,
)
ACTIVITY_WORDS = (
    "аренда",
    "мойка",
    "ремонт",
    "топлив",
    "масло",
    "запчаст",
    "перевоз",
    "диагност",
    "теплов",
    "водоотвед",
    "хвс",
    "вывоз",
    "мусор",
    "тко",
    "энерг",
    "связ",
    "мобайл",
    "картридж",
    "мфу",
    "колес",
    "шин",
    "страх",
    "сотрудник",
    "кап.рем",
    "капитал",
    "соцтакс",
    "автозапчаст",
    "автосервис",
)


def _is_legal_entity(text: str) -> bool:
    t = text.strip()
    return bool(LEGAL_RE.match(t)) or t.upper().startswith("ИП ")


def _looks_like_activity(text: str) -> bool:
    low = text.lower()
    if _is_legal_entity(text):
        return False
    return any(w in low for w in ACTIVITY_WORDS) or len(text) < 60


def _activity_from_name(text: str) -> str:
    hit = _activity_in_parens_anywhere(text)
    if hit:
        return hit
    if not text or _is_legal_entity(text) or len(text) > 45:
        return ""
    low = text.lower()
    if any(w in low for w in ACTIVITY_WORDS):
        return text.strip()
    return ""


def _activity_in_parens_anywhere(rest: str) -> str:
    for m in re.finditer(r"\(([^)]+)\)", rest):
        inner = m.group(1).strip()
        if _looks_like_activity(inner):
            return inner
    return ""


def _contract_hint(folder: Path) -> str:
    hints: list[str] = []
    for p in sorted(folder.iterdir()):
        if not p.is_file():
            continue
        n = p.name.lower()
        if "договор" in n or "соглашен" in n or "контракт" in n:
            hints.append(p.stem[:120])
        if len(hints) >= 2:
            break
    return "; ".join(hints)


def parse_folder(folder: Path) -> dict:
    name = folder.name
    m = NUM_RE.match(name)
    if m:
        num, rest = m.group(1), m.group(2).strip()
    else:
        num, rest = "—", name

    supplier = rest
    activity = ""
    pm = PARENS_RE.search(rest)
    if pm:
        inner = pm.group(1).strip()
        before = rest[: pm.start()].strip()
        if _is_legal_entity(inner):
            supplier = f"{before} ({inner})".strip(" ()") if before else inner
            activity = _activity_from_name(before)
        elif _looks_like_activity(inner):
            supplier = before or rest
            activity = inner
        else:
            supplier = before or rest
            activity = inner

    if not activity:
        activity = _activity_in_parens_anywhere(rest)
    if not activity:
        activity = _activity_from_name(rest)

    hint = _contract_hint(folder)
    note = ""
    if hint and not activity:
        activity = hint
        note = "предмет по имени файла договора"
    elif hint and activity:
        note = f"договор: {hint[:80]}"

    files = sum(1 for _ in folder.rglob("*") if _.is_file())

    return {
        "№_папки": num,
        "название_папки": name,
        "наименование_поставщика": supplier,
        "вид_деятельности_предмет_договора": activity,
        "файлов_в_папке": files,
        "примечание": note,
        "_sort_num": int(num) if num.isdigit() else 9999,
    }


def collect() -> pd.DataFrame:
    rows = [parse_folder(d) for d in SUPPLIERS_ROOT.iterdir() if d.is_dir()]
    df = pd.DataFrame(rows)
    df = df.sort_values(["_sort_num", "название_папки"]).reset_index(drop=True)
    df.insert(0, "№_п_п", range(1, len(df) + 1))
    return df.drop(columns=["_sort_num"])


def build() -> Path:
    if not SUPPLIERS_ROOT.is_dir():
        raise FileNotFoundError(SUPPLIERS_ROOT)

    df = collect()
    out_git = ensure_registers_dir() / REGISTER_FILES[APP_ID]

    desc = pd.DataFrame(
        {
            "пояснение": [
                "Сводная ведомость контрагентов (поставщиков) РКООИ ЦИП «Таганай».",
                f"Источник: {SUPPLIERS_ROOT.relative_to(BASE).as_posix()}/",
                "Папки нумеруются «N. Наименование (предмет)»; внутри — договоры, акты, УПД, оплаты.",
                "Вид деятельности / предмет договора — по названию папки и при необходимости по имени файла договора.",
                f"Записей: {len(df)}; файлов в папках: {int(df['файлов_в_папке'].sum())}.",
                "Исходные PDF — на Яндекс.Диске (прил. 11.1); в GitHub только эта сводная.",
            ]
        }
    )
    totals = pd.DataFrame(
        [
            ("папок_поставщиков", len(df)),
            ("файлов_всего", int(df["файлов_в_папке"].sum())),
            ("дата_сборки", datetime.now().strftime("%d.%m.%Y %H:%M")),
        ],
        columns=["показатель", "значение"],
    )

    with pd.ExcelWriter(out_git, engine="openpyxl") as w:
        desc.to_excel(w, sheet_name="Описание", index=False)
        df.to_excel(w, sheet_name="Поставщики", index=False)
        totals.to_excel(w, sheet_name="Итоги", index=False)

    copy_register_to_case(out_git, "11_Поставщики")
    return out_git


def update_katalog(fname: str) -> None:
    if not KATALOG_PATH.is_file():
        return
    text = KATALOG_PATH.read_text(encoding="utf-8")
    if "11.2 |" in text:
        return
    line = (
        f"| 11.2 | `{fname}` | {datetime.now():%Y-%m-%d} | "
        f"в GitHub: `data/delo_registers/`; PDF — прил. 11.1 (Яндекс.Диск) |"
    )
    marker = "| 11.1 |"
    if marker in text:
        idx = text.find(marker)
        end = text.find("\n", idx)
        text = text[: end + 1] + line + "\n" + text[end + 1 :]
    else:
        text = text.rstrip() + "\n" + line + "\n"
    if "`docs/Поставщики`" not in text and "Разнорядки" in text:
        text = text.replace(
            "| `docs/Разнорядки 2024/`, `docs/Разнорядки 2025/` |",
            "| `docs/Поставщики/` | `11_Поставщики/` (архив `11.1`, свод `11.2`) |\n"
            "| `docs/Разнорядки 2024/`, `docs/Разнорядки 2025/` |",
        )
    KATALOG_PATH.write_text(text, encoding="utf-8")


def main() -> None:
    out = build()
    wb = load_workbook(out, read_only=True)
    n = len(wb.sheetnames)
    wb.close()
    update_register_opis_row(
        APP_ID,
        "Сводная ведомость поставщиков (реестр папок и предмет договоров)",
        n,
    )
    update_katalog(out.name)
    print(f"OK (git): {out} ({n} листов)")


if __name__ == "__main__":
    main()
