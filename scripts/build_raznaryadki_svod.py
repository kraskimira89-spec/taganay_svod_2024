# -*- coding: utf-8 -*-
"""
Сводные ведомости по папкам docs/Разнорядки {2024|2025}.

Разнарядка — план распределения заявок на следующий день между экипажами
(время, адреса; иногда — обратная поездка). Исходники — PDF-сканы; текст
из PDF не извлекается (скан без OCR). Ведомость = реестр файлов + своды
по месяцам и направлениям (Салехард, Лабытнанги и т.д.).

Запуск:
  python scripts/build_raznaryadki_svod.py
  python scripts/build_raznaryadki_svod.py --year 2025
"""

from __future__ import annotations

import argparse
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from pypdf import PdfReader

from delo_github_policy import (
    copy_register_to_case,
    ensure_registers_dir,
    update_register_opis_row,
    REGISTER_FILES,
)

BASE = Path(__file__).resolve().parent.parent
DOCS = BASE / "docs"
CASE_ROOT = BASE / "Дело3а-78-2026Таганай"
KATALOG_PATH = CASE_ROOT / "Каталог_приложений.md"

MONTH_BY_NAME = {
    "январ": 1,
    "феврал": 2,
    "март": 3,
    "апрел": 4,
    "май": 5,
    "мая": 5,
    "июн": 6,
    "июл": 7,
    "август": 8,
    "сентябр": 9,
    "октябр": 10,
    "ноябр": 11,
    "декабр": 12,
}

DATE_RE = re.compile(r"^(\d{2})\.(\d{2})\.(\d{4})$", re.I)


def _month_from_folder(name: str) -> int | None:
    low = name.lower()
    for key, num in MONTH_BY_NAME.items():
        if key in low:
            return num
    return None


def _pdf_pages(path: Path) -> int | None:
    try:
        return len(PdfReader(str(path)).pages)
    except Exception:
        return None


def collect_rows(year: int, count_pages: bool = True) -> list[dict]:
    root = DOCS / f"Разнорядки {year}"
    if not root.is_dir():
        raise FileNotFoundError(root)

    rows: list[dict] = []
    for pdf in sorted(root.rglob("*.pdf")):
        rel = pdf.relative_to(root)
        parts = rel.parts
        month_folder = parts[0] if parts else ""
        month_num = _month_from_folder(month_folder)

        crew = ""
        doc_kind = "ежедневная"
        if pdf.name.lower() == "scan.pdf":
            doc_kind = "месячный свод (Scan.pdf)"
            if len(parts) == 2:
                crew = "—"
            elif len(parts) >= 3:
                crew = parts[1]
        elif len(parts) >= 3:
            crew = parts[1]
        elif len(parts) == 2:
            crew = parts[1] if pdf.name.lower() != "scan.pdf" else "—"

        date_val = None
        date_str = ""
        m = DATE_RE.match(pdf.stem)
        if m:
            d, mo, y = map(int, m.groups())
            date_val = datetime(y, mo, d)
            date_str = date_val.strftime("%d.%m.%Y")
            if month_num is None:
                month_num = mo

        pages = _pdf_pages(pdf) if count_pages else None
        rows.append(
            {
                "год": year,
                "месяц_папка": month_folder,
                "месяц_№": month_num,
                "дата": date_str,
                "дата_sort": date_val,
                "направление_экипаж": crew,
                "тип_файла": doc_kind,
                "имя_файла": pdf.name,
                "относительный_путь": rel.as_posix(),
                "листов_pdf": pages,
                "полный_путь": str(pdf),
            }
        )
    return rows


def _summary_pivot(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        return pd.DataFrame()
    sub = df[df["тип_файла"] == "ежедневная"].copy()
    if sub.empty:
        sub = df.copy()
    g = (
        sub.groupby(["месяц_№", "направление_экипаж"], dropna=False)
        .agg(файлов=("имя_файла", "count"), листов_pdf=("листов_pdf", "sum"))
        .reset_index()
    )
    g = g.sort_values(["месяц_№", "направление_экипаж"])
    return g


def _totals(df: pd.DataFrame, year: int) -> pd.DataFrame:
    daily = df[df["тип_файла"] == "ежедневная"]
    scan = df[df["тип_файла"].str.contains("Scan", na=False)]
    crews = sorted({c for c in df["направление_экипаж"].unique() if c and c != "—"})
    rows = [
        ("год", year),
        ("всего_pdf_файлов", len(df)),
        ("ежедневных_разнарядок", len(daily)),
        ("месячных_scan_pdf", len(scan)),
        ("листов_pdf_всего", int(df["листов_pdf"].fillna(0).sum())),
        ("направлений_экипажей", len(crews)),
        ("направления", ", ".join(crews) if crews else "—"),
    ]
    if not daily.empty and daily["дата_sort"].notna().any():
        dmin = daily["дата_sort"].min()
        dmax = daily["дата_sort"].max()
        rows.append(("период_ежедневных_с", dmin.strftime("%d.%m.%Y")))
        rows.append(("период_ежедневных_по", dmax.strftime("%d.%m.%Y")))
    return pd.DataFrame(rows, columns=["показатель", "значение"])


def _description_sheet(year: int) -> pd.DataFrame:
    text = [
        "Сводная ведомость разнарядок (плановое распределение заявок).",
        f"Год: {year}. Источник: docs/Разнорядки {year}/.",
        "Разнарядка составляется на следующий день: время, места прибытия и доставки;",
        "иногда указывается приблизительное время обратной поездки.",
        "Исходные PDF — сканы; машиночитаемого текста в файлах нет.",
        "Лист «Реестр» — полный перечень файлов с путями и числом листов PDF.",
        "Листы «По месяцам» и «Итоги» — агрегаты для навигации и доказательства объёма.",
    ]
    if year == 2024:
        text.append(
            "За 2024 год в архиве — по одному Scan.pdf на календарный месяц "
            "(детализация по дням — внутри месячного скана)."
        )
    else:
        text.append(
            "За 2025 год — отдельные PDF по датам (имя DD.MM.YYYY.pdf) "
            "в папках месяца и направления (Салехард, Лабытнанги и др.)."
        )
    text.append(
        f"Исходные PDF — Яндекс.Диск (прил. 12.{3 if year == 2024 else 4}); "
        "в GitHub — только эта сводная."
    )
    return pd.DataFrame({"пояснение": text})


def build_workbook(year: int, count_pages: bool = True) -> Path:
    rows = collect_rows(year, count_pages=count_pages)
    df = pd.DataFrame(rows)
    if not df.empty:
        show = df.drop(columns=["дата_sort", "полный_путь"], errors="ignore")
    else:
        show = df

    app_id = f"12.{1 if year == 2024 else 2}"
    out_git = ensure_registers_dir() / REGISTER_FILES[app_id]

    with pd.ExcelWriter(out_git, engine="openpyxl") as w:
        _description_sheet(year).to_excel(w, sheet_name="Описание", index=False)
        show.to_excel(w, sheet_name="Реестр", index=False)
        _summary_pivot(df).to_excel(w, sheet_name="По_месяцам", index=False)
        _totals(df, year).to_excel(w, sheet_name="Итоги", index=False)

    copy_register_to_case(out_git, "12_Разнарядки")
    return out_git


def _sheet_count(xlsx: Path) -> int:
    wb = load_workbook(xlsx, read_only=True)
    n = len(wb.sheetnames)
    wb.close()
    return n


def update_katalog(entries: list[tuple[str, str, str]]) -> None:
    if not KATALOG_PATH.is_file():
        return
    text = KATALOG_PATH.read_text(encoding="utf-8")
    if "12.1 |" in text:
        return
    block = "\n".join(
        f"| {addr} | `{fname}` | {datetime.now():%Y-%m-%d} | {note} |"
        for addr, fname, note in entries
    )
    marker = "| 11.1 |"
    if marker in text:
        idx = text.find(marker)
        end = text.find("\n", idx)
        text = text[: end + 1] + block + "\n" + text[end + 1 :]
    else:
        text = text.rstrip() + "\n" + block + "\n"
    if "## Десять разделов" in text and "| 12 |" not in text:
        text = text.replace(
            "| 11 | `11_Поставщики` |",
            "| 11 | `11_Поставщики` |\n| 12 | `12_Разнарядки` | Плановые разнарядки (распределение заявок по экипажам), сводные ведомости 2024/2025 |",
        )
    KATALOG_PATH.write_text(text, encoding="utf-8")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--year", type=int, choices=(2024, 2025), default=0)
    ap.add_argument("--no-page-count", action="store_true")
    args = ap.parse_args()
    years = [args.year] if args.year else [2024, 2025]

    built: list[Path] = []
    for y in years:
        p = build_workbook(y, count_pages=not args.no_page_count)
        built.append(p)
        print(f"OK {y}: {p}")

    opis_rows = []
    kat_rows = []
    for p in built:
        y = 2024 if "2024" in p.name else 2025
        n = _sheet_count(p)
        app = f"12.{1 if y == 2024 else 2}"
        fname = REGISTER_FILES[app]
        title = f"Сводная ведомость разнарядок соцтакси за {y} год"
        opis_rows.append((app, title, fname, n))
        kat_rows.append(
            (
                app,
                fname,
                f"GitHub: data/delo_registers/; PDF — прил. 12.{3 if y == 2024 else 4} (Яндекс.Диск)",
            )
        )

    for app, title, fname, n in opis_rows:
        update_register_opis_row(app, title, n)
    update_katalog(kat_rows)
    print("Опись и каталог обновлены.")


if __name__ == "__main__":
    main()
