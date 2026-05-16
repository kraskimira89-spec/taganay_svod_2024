# -*- coding: utf-8 -*-
r"""
Ведомость приложений для одного раздела папки дела (формат Excel).

Использование:
  python scripts/sozdat_vedomost_razdela.py "Путь\\к\\01_Иск"
  python scripts/sozdat_vedomost_razdela.py . --yandex-base "https://disk.yandex.ru/d/..."

Колонка «Наименование документа» заполняется эвристикой по имени файла (можно править вручную).

Пакетный запуск по 11 разделам: Sozdat_vse_vedomosti.bat в корне дела.
"""

from __future__ import annotations

import argparse
import re
import subprocess
import sys
from datetime import datetime
from pathlib import Path

# ----- зависимости -----
try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "-q"])
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ARIAL = "Arial"
C_BRAND = "01696F"
# См. sozdat_vedomost_pdf: в Excel — только dd.mm.yyyy (строчные).
XLSX_DATE_SHORT = "dd.mm.yyyy"
C_RED = "A12C7B"
C_GREEN = "E8F1E3"
C_YELLOW = "FFF6CC"  # «Наименование документа» — как в образце (ручная правка)
C_ROSE = "F5E3EC"

THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill(start_color=C_BRAND, end_color=C_BRAND, fill_type="solid")
HEADER_FONT = Font(name=ARIAL, size=11, bold=True, color="FFFFFF")

SECTIONS = {
    "01": {"name": "Исковое заявление", "icon": "📄", "desc": "Иск, ходатайства, отзывы, процессуальные документы"},
    "02": {"name": "Отчёты экспертные", "icon": "📊", "desc": "Обоснования и отчёты по расходам, аналитика"},
    "03": {"name": "Учредительные документы", "icon": "🏢", "desc": "Устав, ЕГРЮЛ, ИНН, протоколы о полномочиях"},
    "04": {"name": "Соглашения с ДСЗН ЯНАО", "icon": "📝", "desc": "Соглашение № 22-25 и допсоглашения"},
    "05": {"name": "Отчёты ЕКЖЯ 2025", "icon": "📅", "desc": "Отчёты системы ЕКЖЯ «Морошка»"},
    "06": {"name": "Финансовые документы", "icon": "💰", "desc": "ОСВ, банк, реестры, сверки"},
    "07": {"name": "Кадровые документы", "icon": "👥", "desc": "Штатное расписание, договоры, приказы"},
    "08": {"name": "Договоры аренды", "icon": "🚐", "desc": "Транспорт, гараж, помещения"},
    "09": {"name": "Нормативные акты", "icon": "⚖️", "desc": "ФЗ, постановления РФ и ЯНАО"},
    "10": {"name": "Рыночные зарплаты", "icon": "💼", "desc": "Подтверждения уровня зарплат в ЯНАО"},
    "11": {"name": "Поставщики (агрегат)", "icon": "📦", "desc": "Сводка по контрагентам"},
}

CASE_NUM = "3а-78/2026"
CASE_COURT = "Суд ЯНАО"
CASE_PARTY = "РКООИ ЦИП «Таганай» (ИНН 8905998693)"


def detect_section(folder_name: str) -> tuple[str, dict]:
    fname = folder_name.lower()
    for sec_num, info in SECTIONS.items():
        if fname.startswith(sec_num) or fname.startswith(f"{sec_num}_"):
            return sec_num, info
    keywords = {
        "01": ["иск", "заявл"],
        "02": ["отчёт", "отчет", "обоснован"],
        "03": ["учредит", "устав"],
        "04": ["соглашен"],
        "05": ["екжя", "морошк"],
        "06": ["финанс", "банк", "осв"],
        "07": ["кадр", "штат"],
        "08": ["аренд"],
        "09": ["нормат", "приказ", "постан"],
        "10": ["зарплат", "фот"],
        "11": ["поставщик", "контраген"],
    }
    for sec_num, words in keywords.items():
        if any(w in fname for w in words):
            return sec_num, SECTIONS[sec_num]
    return "??", {"name": folder_name, "icon": "📁", "desc": "Уточните раздел вручную"}


def suggest_document_title(filename: str) -> str:
    """
    Юридически осмысленное наименование по имени файла / папки (подсказка).
    """
    is_dir = filename.endswith("/")
    raw = filename.rstrip("/")
    path = Path(raw)
    ext = path.suffix.lower()
    stem = path.stem if ext else raw

    stem = re.sub(
        r"^\d{1,2}[._]\d{1,2}[a-zа-яё]?[_\s\-.]+",
        "",
        stem,
        flags=re.IGNORECASE,
    )
    stem = re.sub(r"^\d{1,2}[_\s\-.]+", "", stem)
    text = stem.replace("_", " ").strip()
    text = re.sub(r"\s+", " ", text)
    text = text.replace(" - ", " — ")
    # номера вида 22-25, 20-26; не дублируем уже имеющийся знак "№"
    text = re.sub(r"(?<!№)(?<!№ )\b(\d{2}\s*-\s*\d{2})\b", r"№ \1", text)

    if is_dir:
        if text:
            return f"Комплект документов: {text}"
        return "Комплект документов (подпапка)"

    if not text:
        return path.name

    low = text.lower()
    formal: list[tuple[str, str]] = [
        ("административн", "Административное исковое заявление"),
        ("исков", "Исковое заявление"),
        ("ходатайств", "Ходатайство"),
        ("отзыв", "Отзыв"),
        ("определен", "Определение суда"),
        ("дополнительн", "Дополнение к исковому заявлению"),
        ("допсоглаш", "Дополнительное соглашение"),
        ("соглашен", "Соглашение"),
        ("договор", "Договор"),
        ("гпх", "Договор ГПХ"),
        ("акт свер", "Акт сверки"),
        ("акт ", "Акт "),
        ("протокол", "Протокол"),
        ("приказ", "Приказ"),
        ("постановлен", "Постановление"),
        ("письм", "Письмо"),
        ("уведомлен", "Уведомление"),
        ("доверен", "Доверенность"),
        ("устав", "Устав"),
        ("выписк", "Выписка"),
        ("отчет", "Отчёт"),
        ("отчёт", "Отчёт"),
        ("счет", "Счёт"),
        (
            "счёт",
            "Счёт",
        ),
        ("платеж", "Платёжное поручение"),
        ("квитанц", "Квитанция"),
    ]
    for key, label in formal:
        if key in low:
            if text.lower().startswith(label.lower()[: min(8, len(label))]):
                out = text[0].upper() + text[1:] if text else label
                return out
            if label.endswith(" ") and any(
                text.lower().startswith(x) for x in ("акт ", "письмо ", "отчёт ")
            ):
                return text[0].upper() + text[1:] if text else label.strip()
            return f"{label} ({text})" if len(text) < 80 else f"{label}: {text}"

    if text[0].islower():
        text = text[0].upper() + text[1:]
    return text


def scan_folder(folder_path: Path) -> tuple[Path, list[dict]]:
    folder = folder_path.resolve()
    if not folder.is_dir():
        print(f"Папка не найдена: {folder}")
        sys.exit(1)

    files: list[dict] = []
    for item in sorted(folder.iterdir(), key=lambda x: str(x).lower()):
        if item.name.startswith(".") or item.name.startswith("~"):
            continue
        if "Ведомость" in item.name or "vedomost" in item.name.lower():
            continue
        if item.name.startswith("00_"):
            continue

        if item.is_file():
            if item.name.lower() == "readme.md":
                continue
            if item.name.lower() == "сводная_ведомость.xlsx":
                continue
            stat = item.stat()
            files.append(
                {
                    "name": item.name,
                    "size_kb": round(stat.st_size / 1024, 1),
                    "mtime": datetime.fromtimestamp(stat.st_mtime),
                    "is_dir": False,
                    "full_path": str(item),
                }
            )
        elif item.is_dir():
            sub_only = [f for f in item.rglob("*") if f.is_file()]
            total_size = sum(f.stat().st_size for f in sub_only) / 1024
            files.append(
                {
                    "name": item.name + "/",
                    "size_kb": round(total_size, 1),
                    "mtime": datetime.fromtimestamp(item.stat().st_mtime),
                    "is_dir": True,
                    "sub_count": len(sub_only),
                    "full_path": str(item),
                }
            )
    return folder, files


def extract_appendix_num(filename: str, section_num: str) -> str:
    m = re.match(r"^(\d{1,2}[._]\d{1,2}[a-zа-яё]?)[_.\s-]", filename, re.IGNORECASE)
    if m:
        return m.group(1).replace("_", ".")
    m = re.match(r"^(\d{1,2})[_.\s-]", filename)
    if m and m.group(1) == section_num:
        return f"{section_num}.?"
    return ""


def create_xlsx(
    folder: Path,
    files: list[dict],
    section_num: str,
    section_info: dict,
    output_path: Path,
    yandex_base: str = "",
) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    safe_sheet_section = re.sub(r"[\[\]:*?/\\]", "_", section_num).strip(" ._") or "unknown"
    ws.title = f"Ведомость {safe_sheet_section}"[:31]

    icon = section_info["icon"]
    ws["B2"] = f"{icon}  ВЕДОМОСТЬ ПРИЛОЖЕНИЙ — РАЗДЕЛ {section_num}"
    ws["B2"].font = Font(name=ARIAL, size=16, bold=True, color=C_BRAND)
    ws.merge_cells("B2:H2")
    ws["B3"] = section_info["name"]
    ws["B3"].font = Font(name=ARIAL, size=14, bold=True)
    ws.merge_cells("B3:H3")
    ws["B4"] = section_info["desc"]
    ws["B4"].font = Font(name=ARIAL, size=10, italic=True, color="7A7974")
    ws.merge_cells("B4:H4")
    ws["B5"] = f"Дело № {CASE_NUM} · {CASE_COURT} · {CASE_PARTY}"
    ws["B5"].font = Font(name=ARIAL, size=10, italic=True, color="7A7974")
    ws.merge_cells("B5:H5")

    total_files = len([f for f in files if not f["is_dir"]])
    total_dirs = len([f for f in files if f["is_dir"]])
    total_kb = sum(f["size_kb"] for f in files)
    total_mb = round(total_kb / 1024, 2)

    ws["B7"] = "ВСЕГО ФАЙЛОВ:"
    ws["B7"].font = Font(name=ARIAL, size=11, bold=True, color=C_BRAND)
    ws["B7"].alignment = Alignment(horizontal="right")
    ws["C7"] = total_files
    ws["C7"].font = Font(name=ARIAL, size=11, bold=True)
    if total_dirs > 0:
        ws["D7"] = "ПОДПАПОК:"
        ws["D7"].font = Font(name=ARIAL, size=11, bold=True, color=C_BRAND)
        ws["D7"].alignment = Alignment(horizontal="right")
        ws["E7"] = total_dirs
        ws["E7"].font = Font(name=ARIAL, size=11, bold=True)
    ws["F7"] = "ОБЪЁМ:"
    ws["F7"].font = Font(name=ARIAL, size=11, bold=True, color=C_BRAND)
    ws["F7"].alignment = Alignment(horizontal="right")
    ws["G7"] = f"{total_mb} МБ" if total_mb > 1 else f"{int(total_kb)} КБ"
    ws["G7"].font = Font(name=ARIAL, size=11, bold=True)
    ws["H7"] = datetime.now()
    ws["H7"].number_format = XLSX_DATE_SHORT
    ws["H7"].font = Font(name=ARIAL, size=10, color="7A7974")
    ws["H7"].alignment = Alignment(horizontal="right")

    ws["B9"] = "📂 ОРИГИНАЛЫ ДОКУМЕНТОВ — НА ЯНДЕКС.ДИСКЕ 360:"
    ws["B9"].font = Font(name=ARIAL, size=11, bold=True, color=C_BRAND)
    ws.merge_cells("B9:D9")

    yandex_link = ""
    if yandex_base:
        yandex_link = f"{yandex_base.rstrip('/')}/{folder.name}"

    ws["E9"] = (
        yandex_link if yandex_link else "_________________________________________ (вставьте ссылку)"
    )
    ws["E9"].font = Font(
        name=ARIAL, size=11, color="0563C1", underline="single" if yandex_link else None
    )
    if yandex_link:
        ws["E9"].hyperlink = yandex_link
    ws.merge_cells("E9:H9")

    headers = [
        ("B", "№\nприл.", 12),
        ("C", "Наименование документа", 45),
        ("D", "Имя файла", 38),
        ("E", "Дата", 12),
        ("F", "Размер", 10),
        ("G", "Стр./Файлов", 12),
        ("H", "Ссылка на Я.Диск", 35),
    ]
    for col, title, width in headers:
        cell = ws[f"{col}11"]
        cell.value = title
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[col].width = width
    ws.row_dimensions[11].height = 32
    ws.column_dimensions["A"].width = 3

    for i, f in enumerate(files):
        r = 12 + i
        app_num = extract_appendix_num(f["name"], section_num)
        ws[f"B{r}"] = app_num if app_num else f"{section_num}.{i + 1}"
        ws[f"B{r}"].font = Font(name=ARIAL, size=10, bold=True, color=C_BRAND)
        ws[f"B{r}"].alignment = Alignment(horizontal="center", vertical="center")
        ws[f"B{r}"].border = BORDER
        ws[f"B{r}"].fill = PatternFill(start_color=C_GREEN, end_color=C_GREEN, fill_type="solid")

        suggested = suggest_document_title(f["name"])
        ws[f"C{r}"] = suggested
        ws[f"C{r}"].font = Font(name=ARIAL, size=10)
        ws[f"C{r}"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws[f"C{r}"].border = BORDER
        ws[f"C{r}"].fill = PatternFill(start_color=C_YELLOW, end_color=C_YELLOW, fill_type="solid")

        ws[f"D{r}"] = f["name"]
        ws[f"D{r}"].font = Font(name="Consolas", size=9)
        ws[f"D{r}"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws[f"D{r}"].border = BORDER

        ws[f"E{r}"] = f["mtime"]
        ws[f"E{r}"].number_format = XLSX_DATE_SHORT
        ws[f"E{r}"].font = Font(name=ARIAL, size=10)
        ws[f"E{r}"].alignment = Alignment(horizontal="center", vertical="center")
        ws[f"E{r}"].border = BORDER

        if f["size_kb"] > 1024:
            ws[f"F{r}"] = f"{round(f['size_kb'] / 1024, 1)} МБ"
        else:
            ws[f"F{r}"] = f"{int(f['size_kb'])} КБ"
        ws[f"F{r}"].font = Font(name=ARIAL, size=10)
        ws[f"F{r}"].alignment = Alignment(horizontal="right", vertical="center")
        ws[f"F{r}"].border = BORDER

        if f["is_dir"]:
            ws[f"G{r}"] = f"📁 {f['sub_count']} файлов"
            ws[f"G{r}"].font = Font(name=ARIAL, size=10, italic=True, color="7A7974")
        else:
            ws[f"G{r}"] = ""
        ws[f"G{r}"].alignment = Alignment(horizontal="center", vertical="center")
        ws[f"G{r}"].border = BORDER

        if yandex_base and yandex_link:
            file_link = f"{yandex_link.rstrip('/')}/{f['name'].rstrip('/')}"
            ws[f"H{r}"] = "Открыть →"
            ws[f"H{r}"].hyperlink = file_link
            ws[f"H{r}"].font = Font(name=ARIAL, size=10, color="0563C1", underline="single")
        else:
            ws[f"H{r}"] = "_____________"
            ws[f"H{r}"].font = Font(name=ARIAL, size=9, color="7A7974", italic=True)
        ws[f"H{r}"].alignment = Alignment(horizontal="center", vertical="center")
        ws[f"H{r}"].border = BORDER
        ws.row_dimensions[r].height = 32

    if not files:
        ws["B12"] = (
            "В этой папке пока нет элементов верхнего уровня. "
            "Добавьте файлы и запустите скрипт повторно."
        )
        ws["B12"].font = Font(name=ARIAL, size=11, italic=True, color=C_RED)
        ws.merge_cells("B12:H12")

    last_r = 12 + len(files) + 2
    ws[f"B{last_r}"] = (
        "Председатель правления РКООИ ЦИП «Таганай»: Богдановский С.В.    "
        "____________________"
    )
    ws[f"B{last_r}"].font = Font(name=ARIAL, size=10)
    ws.merge_cells(f"B{last_r}:H{last_r}")

    last_r += 2
    ws[f"B{last_r}"] = "Подсказки:"
    ws[f"B{last_r}"].font = Font(name=ARIAL, size=10, bold=True, color="7A7974")
    tips = [
        "• Колонка «Наименование документа» заполнена автоматически по имени файла — проверьте формулировки.",
        "• Колонка «Стр.» (G) — для PDF укажите число страниц при необходимости.",
        "• Перезапуск скрипта перезапишет файл (в т.ч. наименования).",
    ]
    for tip in tips:
        last_r += 1
        ws[f"B{last_r}"] = tip
        ws[f"B{last_r}"].font = Font(name=ARIAL, size=9, color="7A7974")
        ws.merge_cells(f"B{last_r}:H{last_r}")

    ws.freeze_panes = "A12"
    if files:
        ws.auto_filter.ref = f"B11:H{11 + len(files)}"

    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.3
    ws.page_margins.right = 0.3
    ws.print_title_rows = "1:11"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


def load_yandex_base_default() -> str:
    repo = Path(__file__).resolve().parent.parent
    yml = repo / "docs" / "delo" / "yandex_disk_links.yaml"
    if not yml.is_file():
        return ""
    try:
        import yaml  # type: ignore
    except ImportError:
        return ""
    try:
        data = yaml.safe_load(yml.read_text(encoding="utf-8")) or {}
    except Exception:
        return ""
    archives = data.get("archives") or {}
    return str(archives.get("case_root") or archives.get("delo") or "").strip()


def main() -> None:
    parser = argparse.ArgumentParser(description="Ведомость раздела дела")
    parser.add_argument("folder", nargs="?", default=".", help="Путь к папке раздела")
    parser.add_argument(
        "--yandex-base",
        default="",
        help="Базовая ссылка Я.Диск на корень дела",
    )
    args = parser.parse_args()
    yandex = args.yandex_base.strip() or load_yandex_base_default()

    folder_path = Path(args.folder).resolve()
    print(f"Раздел: {folder_path}")
    section_num, section_info = detect_section(folder_path.name)
    print(f"Номер раздела: {section_num} — {section_info['name']}")

    folder, files = scan_folder(folder_path)
    print(
        f"Файлов (верхний уровень): {len([f for f in files if not f['is_dir']])}, "
        f"подпапок: {len([f for f in files if f['is_dir']])}"
    )

    safe_section_num = re.sub(r'[<>:"/\\\\|?*]', "_", section_num).strip(" ._") or "unknown"
    safe_name = re.sub(r'[<>:"/\\\\|?*]', "_", section_info["name"].replace(" ", "_"))
    output_name = f"00_Ведомость_{safe_section_num}_{safe_name}.xlsx"
    output_path = folder / output_name

    create_xlsx(folder, files, section_num, section_info, output_path, yandex_base=yandex)
    print(f"Готово: {output_path} ({output_path.stat().st_size // 1024} КБ)")


if __name__ == "__main__":
    main()
